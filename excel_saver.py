# excel_saver.py
import openpyxl
import os
import re

def find_excel_file(work_folder):
    """Ищет файл .xlsx в рабочей папке, возвращает путь или None."""
    try:
        for f in os.listdir(work_folder):
            if f.endswith('.xlsx') and not f.startswith('~'):
                return os.path.join(work_folder, f)
    except Exception:
        pass
    return None

def open_workbook(filepath):
    """Открывает workbook, возвращает объект."""
    return openpyxl.load_workbook(filepath, data_only=False)

def save_workbook(workbook, filepath):
    """Сохраняет workbook."""
    workbook.save(filepath)

def get_all_sheets(workbook):
    """Возвращает список всех листов в книге."""
    return workbook.sheetnames

def find_variant_sheet(workbook, variant_number):
    """
    Ищет лист, в названии которого есть номер варианта.
    Возвращает объект листа или None.
    """
    variant_str = str(variant_number)
    for sheet in workbook.worksheets:
        if variant_str in sheet.title:
            return sheet
    return None

def get_sheet_by_name(workbook, sheet_name):
    """Возвращает лист по имени."""
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name]
    return None

def get_all_students(sheet):
    """
    Возвращает список всех учеников (значения в столбце A, начиная с 3 строки).
    """
    students = []
    for row in range(3, sheet.max_row + 1):
        cell_value = sheet.cell(row=row, column=1).value
        if cell_value:
            cell_str = str(cell_value).strip()
            # Пропускаем служебные строки
            if not cell_str.startswith('min:') and not cell_str.startswith('max:') and cell_str != '-':
                students.append(cell_str)
    return students

def find_student_row(sheet, student_name):
    """
    Ищет строку ученика по точному или частичному совпадению.
    Возвращает номер строки или None.
    """
    student_name = student_name.strip().lower()
    for row in range(3, sheet.max_row + 1):
        cell_value = sheet.cell(row=row, column=1).value
        if cell_value:
            cell_str = str(cell_value).strip().lower()
            # Проверяем совпадение (по фамилии или полному имени)
            if student_name in cell_str or cell_str in student_name:
                return row
    return None

def write_scores_to_sheet(sheet, row_idx, scores):
    """
    Записывает баллы в строку row_idx, начиная с колонки E (задание 1) и далее.
    scores: список из 27 чисел (баллы за задания 1..27).
    """
    start_col = 5  # E = 5
    for i, points in enumerate(scores):
        col = start_col + i
        if col <= sheet.max_column:
            sheet.cell(row=row_idx, column=col, value=points)

def get_scores_from_sheet(sheet, row_idx):
    """Считывает баллы из строки (для отладки)."""
    scores = []
    for col in range(5, 5+27):
        val = sheet.cell(row=row_idx, column=col).value
        scores.append(val if isinstance(val, (int, float)) else 0)
    return scores